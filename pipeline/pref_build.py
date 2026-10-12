"""都道府県の産業連関表(2020年・統合大分類37部門ベース)を共通形式に整え、data/pref_37.json を作る。
入力: pref_raw/ (各県の公表ファイル。出典は pref_raw/SOURCES.json)。単位は百万円。"""

import csv, json, pathlib
import openpyxl, pandas as pd

R = pathlib.Path(__file__).resolve().parent.parent
RAW = R / "pref_raw"
OUT = R / "data"
SRC = json.loads((RAW / "SOURCES.json").read_text(encoding="utf-8"))
num = lambda x: float(str(x).replace(",", "")) if x not in (None, "") else 0.0

import re

norm = lambda v: re.sub(r"\s+", "", str(v))


def table(rows, name_row, first, code_col, name_col):
    cols = {}
    for j, v in enumerate(rows[name_row]):
        if v not in (None, ""):
            # 宮城県の「県内需要合計③＝①＋②」のような丸数字以降の注記は外す
            cols.setdefault(re.sub(r"[①-⑳].*$", "", norm(v)), j)  # 同名は最初の列
    rl = [
        (norm(r[code_col]) if r[code_col] is not None else "", norm(r[name_col]), i)
        for i, r in enumerate(rows)
        if i >= first and r[name_col]
    ]
    return cols, rl


def load(slug):
    s = SRC[slug]
    f = RAW / s["file"]
    if f.suffix == ".csv":
        rows = list(csv.reader(open(f, encoding="cp932")))
        name_row, first, code_col, name_col = 1, 2, 0, 1
    else:
        ws = openpyxl.load_workbook(f, read_only=True, data_only=True)[s["sheet"]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        name_row, first, code_col, name_col = s["name_row"], s["first"], s["code_col"], s["name_col"]
        if s.get(
            "blocks"
        ):  # 宮崎県: 表が上下2段に分かれている。[見出し行, 最初の行, 最後の行] ごとに横につなぐ
            b0 = s["blocks"][0]
            out = [[None, None] + [], *[[r[code_col], r[name_col]] for r in rows[b0[1] : b0[2] + 1]]]
            for nr, fr, lr in s["blocks"]:
                width = max(j for j, v in enumerate(rows[nr]) if v not in (None, "")) + 1
                out[0] += rows[nr][name_col + 1 : width]
                block = rows[fr : lr + 1]
                assert [norm(r[name_col]) for r in block] == [norm(r[1]) for r in out[1:]], (
                    slug,
                    "段の行が一致しない",
                )
                for k, r in enumerate(block):
                    out[k + 1] += r[name_col + 1 : width]
            rows, name_row, first, code_col, name_col = out, 0, 1, 0, 1
    k = s.get("scale")
    if k:  # 単位を百万円にそろえる(例: 万円表記なら 0.01)
        rows = [
            r
            if i < first
            else [
                v * k if isinstance(v, (int, float)) and not isinstance(v, bool) and j > code_col + 1 else v
                for j, v in enumerate(r)
            ]
            for i, r in enumerate(rows)
        ]
    cols, rl = table(rows, name_row, first, code_col, name_col)
    return rows, cols, rl, s


PICK = {
    "X": ["生産額", "県内生産額", "府内生産額", "都内生産額", "国内生産額", "県内生産額"],
    "dem": ["都内需要合計", "県内需要合計", "府内需要合計", "県内需要計", "県内需用合計"],
    "int": ["内生部門計"],
    "mexp": ["移輸出計", "輸移出計", "移輸出"],
    "explicit_imp": ["（控除）移輸入計", "移輸入計", "（控除）移輸入"],
}


def col(cols, k):
    for n in PICK[k]:
        if n in cols:
            return cols[n]
    raise KeyError((k, list(cols)[:60]))


def colopt(cols, k):
    for n in PICK[k]:
        if n in cols:
            return cols[n]


def build_okinawa(slug):
    """沖縄県: 全国の基本分類に対応する397×361部門表を、全国の37部門に合算する。"""
    s = SRC[slug]
    ws = openpyxl.load_workbook(RAW / s["file"], read_only=True, data_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    nd = pd.read_csv(OUT / "nodes.csv", encoding="utf-8-sig", dtype={"code": str})
    n108 = nd[nd.level == 108].set_index("code")
    n37 = nd[nd.level == 37].set_index("code")
    to37 = lambda code3: "%02d" % int(n108.loc[code3, "parent37"])
    hdr = [norm(x) for x in rows[5]]
    cd = rows[4]
    ci = {h: j for j, h in reversed(list(enumerate(hdr))) if h}
    first_col = 2
    last_col = [j for j, h in enumerate(hdr) if h == "内生部門計"][0]  # 内生の列は first_col..last_col-1
    for k in ("県内需要合計", "輸出計", "移出計", "県内生産額"):
        assert k in ci, k
    agg = {}
    a = lambda k: agg.setdefault(k, dict(out=0.0, dem=0.0, exp=0.0, va=0.0, intin=0.0))
    ri_int = [i for i, r in enumerate(rows) if i >= 6 and r[1] and norm(r[1]) == "内生部門計"][0]
    ri_va = [i for i, r in enumerate(rows) if i >= 6 and r[1] and norm(r[1]) == "粗付加価値部門計"][0]
    for i in range(6, ri_int):  # 行: 内生部門
        r = rows[i]
        code = str(r[0])
        if not code[:4].isdigit() or int(code[:2]) > 69:
            continue
        k = to37(code[:3])
        g = a(k)
        g["out"] += num(r[ci["県内生産額"]])
        g["dem"] += num(r[ci["県内需要合計"]])
        g["exp"] += num(r[ci["輸出計"]]) + num(r[ci["移出計"]])
    for j in range(first_col, last_col):  # 列: 内生部門(付加価値・中間投入)
        code = str(cd[j])
        if not code[:4].isdigit() or int(code[:2]) > 69:
            continue
        k = to37(code[:3])
        g = a(k)
        g["va"] += num(rows[ri_va][j])
        g["intin"] += num(rows[ri_int][j])
    out = []
    for k, g in sorted(agg.items()):
        mi = g["dem"] + g["exp"] - g["out"]
        assert abs(g["intin"] + g["va"] - g["out"]) <= max(3, g["out"] * 1e-4), (
            slug,
            k,
            "費用側",
            g["intin"] + g["va"],
            g["out"],
        )
        assert mi >= -max(3, g["out"] * 1e-4), (slug, k, "移輸入が負", mi)
        out.append(
            dict(
                c=k,
                n=n37.loc[k, "name"],
                out=g["out"],
                va=g["va"],
                exp=g["exp"],
                imp=max(mi, 0),
                dem=g["dem"],
            )
        )
    return dict(name=s["name"], year=s["year"], url=s["url"], file=s["file"], sectors=out)


# 北海道の38部門 -> 全国の37部門と同じ粗さにそろえた共通部門 (名称, 含まれる全国37部門のコード)
HK_GROUPS = {
    "01": ("農林漁業", ["01"]),
    "06": ("鉱業", ["06"]),
    "11": ("飲食料品", ["11"]),
    "15": ("繊維製品", ["15"]),
    "16": ("木材・家具・パルプ・紙", ["16"]),
    "20": ("化学製品", ["20"]),
    "21": ("石油・石炭製品", ["21"]),
    "22+39": ("プラ・ゴム・印刷等の製造", ["22", "39"]),
    "25": ("窯業・土石製品", ["25"]),
    "26": ("鉄鋼", ["26"]),
    "27": ("非鉄金属", ["27"]),
    "28": ("金属製品", ["28"]),
    "29-35": ("機械（7部門合算）", ["29", "30", "31", "32", "33", "34", "35"]),
    "41": ("建設", ["41"]),
    "46+47": ("電気・ガス・熱供給・水道", ["46", "47"]),
    "48": ("廃棄物処理", ["48"]),
    "51": ("商業", ["51"]),
    "53+55": ("金融・保険・不動産", ["53", "55"]),
    "57": ("運輸・郵便", ["57"]),
    "59": ("情報通信", ["59"]),
    "61": ("公務", ["61"]),
    "63": ("教育・研究", ["63"]),
    "64": ("医療・福祉", ["64"]),
    "65": ("他に分類されない会員制団体", ["65"]),
    "66": ("対事業所サービス", ["66"]),
    "67": ("対個人サービス", ["67"]),
    "68": ("事務用品", ["68"]),
    "69": ("分類不明", ["69"]),
}
HK_MAP = {
    1: "01",
    2: "01",
    3: "01",
    4: "01",
    5: "06",
    6: "11",
    7: "11",
    8: "11",
    9: "15",
    10: "16",
    11: "16",
    12: "22+39",
    13: "20",
    14: "21",
    15: "22+39",
    16: "22+39",
    17: "25",
    18: "26",
    19: "26",
    20: "27",
    21: "28",
    22: "29-35",
    23: "22+39",
    24: "41",
    25: "46+47",
    26: "48",
    27: "51",
    28: "53+55",
    29: "57",
    30: "59",
    31: "61",
    32: "63",
    33: "64",
    34: "65",
    35: "66",
    36: "67",
    37: "68",
    38: "69",
}


def build_hokkaido(slug):
    s = SRC[slug]
    ws = (
        openpyxl.load_workbook(RAW / s["file"], read_only=True, data_only=True)[s["sheet"]]
        if s.get("sheet")
        else openpyxl.load_workbook(RAW / s["file"], read_only=True, data_only=True).worksheets[0]
    )
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    assert (
        norm(rows[7][1]) == "耕種農業"
        and norm(rows[44][1]) == "分類不明"
        and norm(rows[45][1]) == "内生部門計"
        and norm(rows[53][1]) == "粗付加価値部門計"
        and "生産額" in norm(rows[54][1])
    )
    assert (
        str(rows[3][40]) == "38" and norm(rows[7][0]) == "01"
    )  # 部門コード列の位置の確認(列3〜40が38部門、41が内生部門計)
    assert "".join(norm(rows[r][61]) for r in range(3, 7)).endswith("生産額")
    COL = dict(int=41, dem=50, exp=53, X=61, imp=59)
    g = {}
    for k in range(38):
        i, j = 7 + k, 3 + k
        r = rows[i]
        X = num(r[COL["X"]])
        dem = num(r[COL["dem"]])
        e = num(r[COL["exp"]])
        mi = dem + e - X
        va = num(rows[53][j])
        intin = num(rows[45][j])
        assert abs(intin + va - X) <= max(2, X * 1e-6), (slug, k + 1, "費用側", intin + va, X)
        assert abs(abs(num(r[COL["imp"]])) - mi) <= max(2, X * 1e-6), (
            slug,
            k + 1,
            "移輸入",
            num(r[COL["imp"]]),
            mi,
        )
        a = g.setdefault(HK_MAP[k + 1], dict(out=0, va=0, exp=0, imp=0, dem=0))
        for f, v in (("out", X), ("va", va), ("exp", e), ("imp", mi), ("dem", dem)):
            a[f] += v
    out = [dict(c=key, n=HK_GROUPS[key][0], nat_codes=HK_GROUPS[key][1], **v) for key, v in g.items()]
    return dict(name=s["name"], year=s["year"], url=s["url"], file=s["file"], sectors=out, coarse=True)


def nat37():
    nd = pd.read_csv(OUT / "nodes.csv", encoding="utf-8-sig", dtype={"code": str})
    return {norm(n): c for c, n in zip(nd[nd.level == 37].code, nd[nd.level == 37].name)}


def build_by_name(slug):
    """部門の分け方が全国の37部門と違う県(宮城38・大分36・広島45・秋田39など)。
    部門名で全国の37部門に対応づけて合算する。SOURCES の by_name で名前ごとの行き先(全国コード、
    または groups のキー)を指定し、指定のない名前は全国の部門名と同じものに対応づける。
    全国の2部門以上をまとめた県の部門(例: 鉄鋼・非鉄金属)は groups に名前と全国コードを書く。"""
    rows, cols, rl, s = load(slug)
    NAT = nat37()
    to = {norm(k): v for k, v in s.get("by_name", {}).items()}
    groups = s.get("groups", {})
    ti = [i for c, n, i in rl if n in ("内生部門計", "内生部門（中間投入）計")][0]
    end = ti
    VA_r = [i for c, n, i in rl if n == "粗付加価値部門計"][0]
    cX, cD = col(cols, "X"), col(cols, "dem")
    cME = colopt(cols, "mexp")
    cEI = colopt(cols, "explicit_imp")
    g = {}
    for c, n, i in rl:
        if i >= end:
            break
        key = to.get(n) or NAT.get(n)
        assert key, (slug, "対応先のない部門", n)
        r = rows[i]
        j = cols.get(n)
        X = num(r[cX])
        e = (
            num(r[cME])
            if cME is not None
            else num(r[cols["輸出" if "輸出" in cols else "輸出計"]]) + num(r[cols["移出"]])
        )
        dem = num(r[cD])
        mi = dem + e - X
        va, intin = num(rows[VA_r][j]), num(rows[ti][j])
        if cEI is not None:
            assert abs(abs(num(r[cEI])) - mi) <= max(2, X * 1e-6), (slug, n, "移輸入", num(r[cEI]), mi)
        assert abs(intin + va - X) <= max(2, X * 1e-6), (slug, n, "費用側", intin + va, X)
        a = g.setdefault(key, dict(out=0, va=0, exp=0, imp=0, dem=0))
        for f, v in (("out", X), ("va", va), ("exp", e), ("imp", max(mi, 0)), ("dem", dem)):
            a[f] += v
    names = {v: k for k, v in NAT.items()}
    nd = pd.read_csv(OUT / "nodes.csv", encoding="utf-8-sig", dtype={"code": str})
    nat_name = dict(zip(nd[nd.level == 37].code, nd[nd.level == 37].name))
    out = []
    for key, v in g.items():
        if key in groups:
            out.append(dict(c=key, n=groups[key][0], nat_codes=groups[key][1], **v))
        else:
            assert key in names, (slug, key)
            out.append(dict(c=key, n=nat_name[key], **v))
    return dict(name=s["name"], year=s["year"], url=s["url"], file=s["file"], sectors=out)


def build(slug):
    if SRC[slug].get("special") == "by_name":
        return build_by_name(slug)
    if SRC[slug].get("special") == "okinawa":
        return build_okinawa(slug)
    if SRC[slug].get("special") == "hokkaido":
        return build_hokkaido(slug)
    rows, cols, rl, s = load(slug)
    sec = [
        (c.lstrip("A").zfill(2), n, i)
        for c, n, i in rl
        if c.lstrip("A").isdigit() and int(c.lstrip("A")) <= 69
    ]
    if s.get(
        "seq37"
    ):  # 福井県・鳥取県: 部門コードが01からの連番。全国の37部門と同じ並び(鳥取は農・林・漁が別の39部門)なので、順番で全国のコードに対応づける
        nd = pd.read_csv(OUT / "nodes.csv", encoding="utf-8-sig", dtype={"code": str})
        nc = list(nd[nd.level == 37].code)
        K = s["seq37"]
        off = K - 37  # K=37 or 39
        cv = lambda n: nc[0] if n <= off + 1 else nc[n - 1 - off]
        sec = [(cv(int(c)), n, i) for c, n, i in sec if int(c) <= K]
        assert len(sec) == K, len(sec)
        chk = [
            (c, n) for c, n, i in sec if norm(nd[nd.code == c].name.iloc[0])[:2] != n[:2] and c != nc[0]
        ]  # 名称の先頭2文字が全国の部門と一致すること
        assert not chk, chk
    sec += [("HQ", n, i) for c, n, i in rl if n == "本社"]  # 東京都のみの部門
    assert any(n in PICK["X"] for c, n, i in rl), (slug, "生産額の行が見つからない")
    VA_r = [i for c, n, i in rl if n == "粗付加価値部門計"][0]
    ti = [i2 for c2, n2, i2 in rl if n2 == "内生部門計"][0]
    cX, cD = col(cols, "X"), col(cols, "dem")
    cME = colopt(cols, "mexp")
    cEI = colopt(cols, "explicit_imp")
    out = []
    for c, n, i in sec:
        r = rows[i]
        j = cols.get(n)
        X = num(r[cX])
        if X <= 0:
            continue
        va = num(rows[VA_r][j]) if j is not None else None
        intin = num(rows[ti][j]) if j is not None else None
        e = (
            num(r[cME])
            if cME is not None
            else num(r[cols["輸出" if "輸出" in cols else "輸出計"]]) + num(r[cols["移出"]])
        )
        dem = num(r[cD])
        mi = dem + e - X  # 移輸入 = 域内需要 + 移輸出 - 生産額 (需給バランスから)
        # 検証: 公表された移輸入の合計列があれば一致すること、費用側(中間投入+付加価値=生産額)が合うこと
        if cEI is not None:
            tax = sum(
                abs(num(r[cols[k]]))
                for k in cols
                if k.startswith("（控除）") and ("関税" in k or "輸入品商品税" in k)
            )  # 鹿児島県は関税・輸入品商品税を別の列に分けている
            assert min(abs(abs(num(r[cEI])) - mi), abs(abs(num(r[cEI])) + tax - mi)) <= max(2, X * 1e-6), (
                slug,
                n,
                "移輸入",
                num(r[cEI]),
                mi,
            )
        if va is not None and intin is not None:
            assert abs(intin + va - X) <= max(2, X * 1e-6), (slug, n, "費用側", intin + va, X)
        assert mi >= -max(2, X * 1e-6), (slug, n, "移輸入が負", mi)
        out.append(dict(c=c, n=n, out=X, va=va, exp=e, imp=max(mi, 0), dem=dem))
    m = s.get("merge")
    if m:  # 兵庫県の39部門を全国の37部門に合わせる(農・林・漁を農林漁業に合算)
        nat = {"01": "農林漁業"}
        g = {}
        for o in out:
            k = m.get(o["c"], o["c"])
            a = g.setdefault(k, dict(c=k, n=nat.get(k, o["n"]), out=0, va=0, exp=0, imp=0, dem=0))
            for f in ("out", "va", "exp", "imp", "dem"):
                a[f] += o[f]
        out = list(g.values())
    return dict(name=s["name"], year=s["year"], url=s["url"], file=s["file"], sectors=out)


if __name__ == "__main__":
    nd = pd.read_csv(OUT / "nodes.csv", encoding="utf-8-sig", dtype={"code": str})
    n37 = nd[nd.level == 37].set_index("code")
    nat = n37.output.astype(float)
    natt = nat.sum()
    res = {}
    for slug in SRC:
        p = build(slug)
        T = sum(o["out"] for o in p["sectors"])
        VA = sum(o["va"] or 0 for o in p["sectors"])
        for o in p["sectors"]:
            o["share"] = o["out"] / T
            codes = o.get("nat_codes", [o["c"]])
            ns = (
                sum(float(nat.get(c, 0)) for c in codes) / natt
                if all(c in nat.index for c in codes)
                else None
            )
            o["nat_share"] = ns
            o["lq"] = (o["share"] / ns) if ns else None
            o["va_ratio"] = (o["va"] / o["out"]) if o["va"] is not None else None
            o["exp_ratio"] = o["exp"] / o["out"]
            o["imp_ratio"] = o["imp"] / o["dem"] if o["dem"] else None
        p["total_out"] = T
        p["total_va"] = VA
        p["exp_total"] = sum(o["exp"] for o in p["sectors"])
        p["imp_total"] = sum(o["imp"] for o in p["sectors"])
        p["note"] = SRC[slug].get("note", "")
        res[slug] = p
        print(slug, p["name"], round(T / 1e6, 1), "兆円", len(p["sectors"]), "部門", "VA", round(VA / 1e6, 1))
    (OUT / "pref_37.json").write_text(json.dumps(res, ensure_ascii=False), encoding="utf-8")
