"""3時点(2011/2015/2020)の産業連関データから、部門別・取引別の時系列を作る。単位は10億円。
- 2011・2015年の取引基本表は単位が100万円なので1/1000して10億円に揃える。
- 接続表(基本分類)を2020年分類の108部門・37部門に集約して、名目値と実質値(2020年価格)の成長を計算する。
入力: data/*.xlsx, kihon_seisan_37.xlsx(2020年37部門), bumonbunrui.xlsx(2020年部門分類表)
出力: out/sector_timeseries_108.csv, sector_timeseries_37.csv, edges37_timeseries.csv, ts_checks.csv
"""
import openpyxl, numpy as np, pandas as pd, os, re
D = os.path.dirname(os.path.abspath(__file__)); OUT = os.path.join(D, "out")
YEARS = (2011, 2015, 2020); CAGR = lambda a, b, n: np.where((a > 0) & (b > 0), (b / np.where(a > 0, a, 1)) ** (1 / n) - 1, np.nan)

# ---------- 37部門の取引基本表(3時点) ----------
FILES37 = {2011: ("data/2011_kihon37.xlsx", 1e-3), 2015: ("data/2015_kihon37.xlsx", 1e-3), 2020: ("kihon_seisan_37.xlsx", 1.0)}
def load37(path, scale):
    ws = openpyxl.load_workbook(os.path.join(D, path), data_only=True).worksheets[0]
    hdr = next(r for r in range(1, 6) if ws.cell(r, 3).value == "01")
    col = {str(ws.cell(hdr, c).value): c for c in range(3, ws.max_column + 1) if ws.cell(hdr, c).value}
    row = {str(ws.cell(r, 1).value): r for r in range(hdr + 2, ws.max_row + 1) if ws.cell(r, 1).value}
    names = {str(ws.cell(hdr, c).value): ws.cell(hdr + 1, c).value for c in range(3, 3 + 37)}
    codes = list(names)[:37]; assert len(codes) == 37 and codes[0] == "01" and codes[-1] == "69"
    g = lambda r, c: float(ws.cell(r, c).value or 0) * scale
    Z = np.array([[g(row[i], col[j]) for j in codes] for i in codes])
    ent = lambda k: np.array([g(row[i], col[k]) for i in codes])
    rowent = lambda k: np.array([g(row[k], col[j]) for j in codes])
    d = dict(codes=codes, Z=Z, X=ent("97"), Xc=rowent("97"),
             fd={"民間消費": ent("72"), "政府消費": ent("73"), "公的投資": ent("74"), "民間投資": ent("75"), "在庫": ent("76"), "輸出": ent("80"), "家計外消費": ent("71")},
             imp=ent("87"), va={"賃金": rowent("91"), "営業余剰": rowent("92"), "減価償却": rowent("93"), "間接税": rowent("94"), "補助金": -rowent("95") * -1})
    d["va"]["補助金"] = rowent("95")           # 表では「(控除)経常補助金」の符号は負のまま
    d["adj"] = ent("77") if "77" in col else np.zeros(37)
    return d
M = {y: load37(*FILES37[y]) for y in YEARS}
codes = M[2020]["codes"]; assert all(M[y]["codes"] == codes for y in YEARS)
names37 = dict(zip(codes, [openpyxl.load_workbook(os.path.join(D, "kihon_seisan_37.xlsx"), data_only=True).worksheets[0].cell(3, 3 + i).value for i in range(37)]))

# ---------- 接続表(基本分類) -> 2020年の108部門・37部門 ----------
wb = openpyxl.load_workbook(os.path.join(D, "bumonbunrui.xlsx"), data_only=True); s = wb.worksheets[0]
b2m, b2b = {}, {}; cur_m = cur_b = None; mid_name = {}; big_name = {}
for r in s.iter_rows(min_row=8, values_only=True):
    if r[7]: cur_m = str(r[7]).strip(); mid_name[cur_m] = str(r[8]).strip() if r[8] else mid_name.get(cur_m, "")
    if r[9]: cur_b = str(r[9]).strip(); big_name[cur_b] = str(r[10]).strip() if r[10] else big_name.get(cur_b, "")
    if r[2] and r[3] and cur_m: key = str(r[2]).strip() + str(r[3]).strip().rstrip("P"); b2m[key] = cur_m; b2b[key] = cur_b   # 末尾のP(屑・副産物等)は接続表のコードでは付かない
cw = openpyxl.load_workbook(os.path.join(D, "data/connect_output_basic.xlsx"), data_only=True).worksheets[0]
rows = []
for r in range(4, cw.max_row + 1):
    code = cw.cell(r, 1).value
    if not code or str(code).startswith("7000"): continue
    code = str(code); rec = dict(code=code, name=str(cw.cell(r, 2).value).strip())
    for k, (nm, c0) in {"n11": ("nominal", 3), "r11": ("real", 6), "n15": ("nominal", 9), "r15": ("real", 12), "n20": ("nominal", 15)}.items():
        rec[k] = float(cw.cell(r, c0).value or 0); rec[k + "_ex"] = float(cw.cell(r, c0 + 1).value or 0); rec[k + "_im"] = float(cw.cell(r, c0 + 2).value or 0)
    rows.append(rec)
C = pd.DataFrame(rows)
C["mid"] = C.code.map(b2m); C["big"] = C.code.map(b2b)
unm = C[C.mid.isna()]; print("接続表の基本分類:", len(C), "行 / 対応づけできず:", len(unm), unm.code.tolist()[:10])
assert C.mid.notna().all(), unm[["code", "name"]].head(10)
C["r20"] = C.n20; C["r20_ex"] = C.n20_ex; C["r20_im"] = C.n20_im
agg = lambda key: C.groupby(key)[[c for c in C.columns if re.match(r"[nr]\d\d", c)]].sum()
A108, A37 = agg("mid"), agg("big")

# ---------- 検算: 接続表(2020年分類)と各年の公式表(当時の分類)の37部門生産額 ----------
chk = []
for y, k in ((2011, "n11"), (2015, "n15"), (2020, "n20")):
    t = pd.Series(M[y]["X"], index=codes); c = A37[k].reindex(codes)
    chk.append(dict(year=y, 公式表の生産額合計=t.sum(), 接続表の生産額合計=c.sum(), 部門別の最大差_10億円=float((t - c).abs().max()), 最大差の部門=names37[(t - c).abs().idxmax()], 最大差の割合=float(((t - c).abs() / t).max())))
chk = pd.DataFrame(chk); print(chk.round(1).to_string(index=False)); chk.to_csv(os.path.join(OUT, "ts_checks.csv"), index=False, encoding="utf-8-sig")

# ---------- 部門別の時系列 ----------
def sector_ts(A, idx, names, extra):
    d = pd.DataFrame({"code": idx, "name": [names[i] for i in idx]})
    for k in ("n11", "r11", "n15", "r15", "n20"):
        d["output_" + k] = A[k].reindex(idx).values; d["export_" + k] = A[k + "_ex"].reindex(idx).values; d["import_" + k] = A[k + "_im"].reindex(idx).values
    d["real_cagr_2011_2015"] = CAGR(d.output_r11.values, d.output_r15.values, 4); d["real_cagr_2015_2020"] = CAGR(d.output_r15.values, d.output_n20.values, 5)
    d["real_cagr_2011_2020"] = CAGR(d.output_r11.values, d.output_n20.values, 9)
    d["price_chg_2011_2020"] = d.output_r11 / d.output_n11 - 1      # 2011年の実質(2020年価格)÷名目 = 2011→2020の価格上昇率
    for k, v in extra.items(): d[k] = v
    return d
mid_idx = sorted(set(C.mid)); big_of_mid = C.drop_duplicates("mid").set_index("mid").big
T108 = sector_ts(A108, mid_idx, mid_name, {"parent37": [big_of_mid[m] for m in mid_idx]})
T108["parent37_name"] = T108.parent37.map(names37)
T37 = sector_ts(A37, codes, names37, {})
# 37部門に、取引基本表からの付加価値・最終需要・輸入の3時点データを追加(10億円)
for y in YEARS:
    m = M[y]; X = m["X"]; va = m["va"]; vsum = va["賃金"] + va["営業余剰"] + va["減価償却"] + va["間接税"] + va["補助金"]
    T37[f"io_output_{y}"] = X; T37[f"io_intermediate_purchases_{y}"] = m["Z"].sum(0); T37[f"io_intermediate_sales_{y}"] = m["Z"].sum(1)
    T37[f"va_ratio_{y}"] = vsum / X; T37[f"wage_ratio_{y}"] = va["賃金"] / X; T37[f"profit_ratio_{y}"] = va["営業余剰"] / X
    T37[f"export_ratio_{y}"] = m["fd"]["輸出"] / X; T37[f"import_ratio_{y}"] = -m["imp"] / (m["Z"].sum(1) + sum(v for k, v in m["fd"].items() if k != "輸出") + m["adj"])
    T37[f"consumption_ratio_{y}"] = m["fd"]["民間消費"] / X; T37[f"gov_ratio_{y}"] = (m["fd"]["政府消費"] + m["fd"]["公的投資"]) / X
T108.round(6).to_csv(os.path.join(OUT, "sector_timeseries_108.csv"), index=False, encoding="utf-8-sig"); T37.round(6).to_csv(os.path.join(OUT, "sector_timeseries_37.csv"), index=False, encoding="utf-8-sig")

# ---------- 取引エッジの時系列(37部門) ----------
price = {y: pd.Series((A37[f"n{y%100}"] / A37[f"r{y%100}"]).reindex(codes).values, index=codes) for y in (2011, 2015)}; price[2020] = pd.Series(1.0, index=codes)
rows = []
for i, ci in enumerate(codes):
    for j, cj in enumerate(codes):
        if i == j: continue
        v = {y: M[y]["Z"][i, j] for y in YEARS}
        if max(v.values()) <= 0: continue
        sh = {y: v[y] / M[y]["X"][j] for y in YEARS}; real = {y: v[y] / price[y][ci] for y in YEARS}
        rows.append(dict(seller=ci, buyer=cj, seller_name=names37[ci], buyer_name=names37[cj], amount_2011=v[2011], amount_2015=v[2015], amount_2020=v[2020],
                         real_2011=real[2011], real_2015=real[2015], real_2020=real[2020], share_of_buyer_cost_2011=sh[2011], share_of_buyer_cost_2015=sh[2015], share_of_buyer_cost_2020=sh[2020],
                         share_of_seller_sales_2020=v[2020] / M[2020]["X"][i]))
E = pd.DataFrame(rows)
E["real_cagr_2011_2015"] = CAGR(E.real_2011.values, E.real_2015.values, 4); E["real_cagr_2015_2020"] = CAGR(E.real_2015.values, E.real_2020.values, 5); E["real_cagr_2011_2020"] = CAGR(E.real_2011.values, E.real_2020.values, 9)
E["share_chg_2011_2020_pt"] = (E.share_of_buyer_cost_2020 - E.share_of_buyer_cost_2011) * 100
E.round(6).to_csv(os.path.join(OUT, "edges37_timeseries.csv"), index=False, encoding="utf-8-sig")
print("108部門", len(T108), "37部門", len(T37), "エッジ(37部門)", len(E))
print("名目の国内生産額合計(公式表, 兆円):", {y: round(M[y]["X"].sum() / 1000, 1) for y in YEARS}); print("実質(2020年価格, 接続表, 兆円):", {y: round(A37[f'r{y%100}'].sum() / 1000, 1) if y < 2020 else round(A37['n20'].sum() / 1000, 1) for y in YEARS})
